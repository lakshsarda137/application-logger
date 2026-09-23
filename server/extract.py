"""Plain-text extraction from uploaded documents (for search and classification)."""

from __future__ import annotations

import io
import logging

log = logging.getLogger(__name__)

MAX_TEXT_CHARS = 500_000


def _kind(filename: str, mime: str | None) -> str:
    name = (filename or "").lower()
    mime = (mime or "").lower()
    if name.endswith(".pdf") or mime == "application/pdf":
        return "pdf"
    if name.endswith(".docx") or "wordprocessingml" in mime:
        return "docx"
    if name.endswith((".txt", ".md")) or mime.startswith("text/"):
        return "text"
    return "unknown"


def extract_text(content: bytes, filename: str, mime: str | None = None) -> str:
    """Best-effort text extraction. Returns "" if the format is unsupported or broken."""
    kind = _kind(filename, mime)
    try:
        if kind == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif kind == "docx":
            import docx

            document = docx.Document(io.BytesIO(content))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        elif kind == "text":
            text = content.decode("utf-8", errors="replace")
        else:
            return ""
    except Exception as e:  # corrupt or encrypted files shouldn't block logging
        log.warning("Could not extract text from %s: %s", filename, e)
        return ""
    return text.replace("\x00", "")[:MAX_TEXT_CHARS]
